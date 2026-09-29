"""
Layout-Aware Section Document Parser for Yojaka AI (Deliverable 17.1).
Replaces arbitrary character-length splitting with semantic layout-aware block parsing.
Extracts PyMuPDF layout blocks (PDF), DOCX paragraph structures, and TXT sections,
preserving section headers, work history roles, and bullet point relationships.
"""

from dataclasses import dataclass, field
import io
import logging
from pathlib import Path
import re
from typing import List, Dict, Any, Optional

logger = logging.getLogger("section_parser")

CANONICAL_SECTION_MAP = {
    "SUMMARY": "summary",
    "PROFESSIONAL SUMMARY": "summary",
    "EXECUTIVE SUMMARY": "summary",
    "OBJECTIVE": "summary",
    "CAREER OBJECTIVE": "summary",
    "EXPERIENCE": "experience",
    "WORK EXPERIENCE": "experience",
    "PROFESSIONAL EXPERIENCE": "experience",
    "WORK HISTORY": "experience",
    "EMPLOYMENT HISTORY": "experience",
    "TECHNICAL SKILLS": "skills",
    "SKILLS": "skills",
    "CORE COMPETENCIES": "skills",
    "AREAS OF EXPERTISE": "skills",
    "EDUCATION": "education",
    "ACADEMIC BACKGROUND": "education",
    "EDUCATIONAL BACKGROUND": "education",
    "PROJECTS": "projects",
    "KEY PROJECTS": "projects",
    "TECHNICAL PROJECTS": "projects",
    "CERTIFICATIONS": "certifications",
    "AWARDS": "certifications",
    "HONORS": "certifications",
    "PUBLICATIONS": "publications",
}


@dataclass
class ParsedSectionChunk:
    """Represents a coherent, semantically bounded document chunk."""

    section_title: str
    section_type: str  # e.g., 'experience', 'skills', 'education', 'summary', 'projects', 'general'
    content: str
    raw_content: str
    role: Optional[str] = None
    company: Optional[str] = None
    duration_years: Optional[float] = None
    skills: List[str] = field(default_factory=list)
    block_index: int = 0
    metadata: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "section_title": self.section_title,
            "section_type": self.section_type,
            "content": self.content,
            "raw_content": self.raw_content,
            "role": self.role,
            "company": self.company,
            "duration_years": self.duration_years,
            "skills": self.skills,
            "block_index": self.block_index,
            "metadata": self.metadata,
        }


class SectionParser:
    """
    Layout-aware section parser leveraging PyMuPDF bounding blocks,
    docx paragraph hierarchies, and text section boundary heuristics.
    """

    def __init__(self, max_chunk_chars: int = 1200):
        self.max_chunk_chars = max_chunk_chars

    @staticmethod
    def _is_section_header(line: str) -> Optional[str]:
        """Detects if a line represents a canonical resume section header."""
        clean = line.strip().strip("#").strip(":").strip("-").strip("*").strip()
        if not clean or len(clean) > 45:
            return None
        upper = clean.upper()
        if upper in CANONICAL_SECTION_MAP:
            return upper
        for header in CANONICAL_SECTION_MAP:
            if upper == header or upper.startswith(f"{header}:") or upper.startswith(f"{header} -"):
                return header
        return None

    def parse_pdf(self, stream_or_path: Any) -> List[ParsedSectionChunk]:
        """Parses PDF layout blocks into structured section chunks via PyMuPDF."""
        try:
            import pymupdf  # type: ignore
        except (ImportError, AttributeError):
            try:
                import fitz as pymupdf  # type: ignore
            except ImportError:
                logger.warning("PyMuPDF not available, falling back to text parsing.")
                return []

        doc = None
        try:
            if isinstance(stream_or_path, (str, Path)):
                doc = pymupdf.open(str(stream_or_path))
            elif isinstance(stream_or_path, bytes):
                doc = pymupdf.open(stream=stream_or_path, filetype="pdf")
            elif hasattr(stream_or_path, "read"):
                bytes_data = stream_or_path.read()
                doc = pymupdf.open(stream=bytes_data, filetype="pdf")
            else:
                return []

            raw_blocks = []
            for page in doc:
                # get_text("blocks") returns (x0, y0, x1, y1, text, block_no, block_type)
                for b in page.get_text("blocks"):
                    if b[6] == 0:  # Text block
                        text = b[4].strip()
                        if text:
                            raw_blocks.append(text)

            return self._structure_blocks(raw_blocks)
        except Exception as e:
            logger.warning(f"PyMuPDF layout block parsing failed: {e}")
            return []
        finally:
            if doc is not None:
                try:
                    doc.close()
                except Exception:
                    pass

    def parse_docx(self, stream_or_path: Any) -> List[ParsedSectionChunk]:
        """Parses DOCX document paragraphs into structured section chunks."""
        try:
            from docx import Document  # type: ignore

            if isinstance(stream_or_path, (str, Path)):
                doc = Document(str(stream_or_path))
            elif isinstance(stream_or_path, bytes):
                doc = Document(io.BytesIO(stream_or_path))
            elif hasattr(stream_or_path, "read"):
                doc = Document(stream_or_path)
            else:
                return []

            blocks = []
            buf = []
            for p in doc.paragraphs:
                txt = p.text.strip()
                if not txt:
                    if buf:
                        blocks.append("\n".join(buf))
                        buf = []
                    continue
                if p.style.name.startswith("Heading") or self._is_section_header(txt):
                    if buf:
                        blocks.append("\n".join(buf))
                        buf = []
                    blocks.append(txt)
                else:
                    buf.append(txt)

            if buf:
                blocks.append("\n".join(buf))

            return self._structure_blocks(blocks)
        except Exception as e:
            logger.warning(f"DOCX layout parsing failed: {e}")
            return []

    def parse_text(self, text: str) -> List[ParsedSectionChunk]:
        """Parses raw text content by section boundaries and double-newline paragraphs."""
        lines = text.splitlines()
        blocks = []
        buf = []

        for line in lines:
            line_str = line.strip()
            if not line_str:
                if buf:
                    blocks.append("\n".join(buf))
                    buf = []
                continue

            if self._is_section_header(line_str):
                if buf:
                    blocks.append("\n".join(buf))
                    buf = []
                blocks.append(line_str)
            else:
                buf.append(line)

        if buf:
            blocks.append("\n".join(buf))

        return self._structure_blocks(blocks)

    def _structure_blocks(self, blocks: List[str]) -> List[ParsedSectionChunk]:
        """
        Groups sequence of text blocks into canonical section chunks,
        preserving role-to-bullet hierarchy.
        """
        chunks: List[ParsedSectionChunk] = []
        current_section = "GENERAL"
        current_type = "general"
        current_role: Optional[str] = None
        current_company: Optional[str] = None
        section_buf: List[str] = []
        block_idx = 0

        def flush_current_chunk():
            nonlocal block_idx
            if not section_buf:
                return
            combined_text = "\n".join(section_buf).strip()
            if not combined_text:
                section_buf.clear()
                return

            # If section text is larger than max_chunk_chars, split gracefully on paragraph breaks
            if len(combined_text) > self.max_chunk_chars and "\n\n" in combined_text:
                sub_parts = combined_text.split("\n\n")
                sub_buf = []
                for sp in sub_parts:
                    sp = sp.strip()
                    if not sp:
                        continue
                    if sum(len(x) for x in sub_buf) + len(sp) > self.max_chunk_chars and sub_buf:
                        part_text = "\n\n".join(sub_buf).strip()
                        chunks.append(
                            ParsedSectionChunk(
                                section_title=current_section,
                                section_type=current_type,
                                content=part_text,
                                raw_content=part_text,
                                role=current_role,
                                company=current_company,
                                block_index=block_idx,
                            )
                        )
                        block_idx += 1
                        sub_buf = [sp]
                    else:
                        sub_buf.append(sp)
                if sub_buf:
                    part_text = "\n\n".join(sub_buf).strip()
                    chunks.append(
                        ParsedSectionChunk(
                            section_title=current_section,
                            section_type=current_type,
                            content=part_text,
                            raw_content=part_text,
                            role=current_role,
                            company=current_company,
                            block_index=block_idx,
                        )
                    )
                    block_idx += 1
            else:
                chunks.append(
                    ParsedSectionChunk(
                        section_title=current_section,
                        section_type=current_type,
                        content=combined_text,
                        raw_content=combined_text,
                        role=current_role,
                        company=current_company,
                        block_index=block_idx,
                    )
                )
                block_idx += 1
            section_buf.clear()

        for b in blocks:
            lines = b.splitlines()
            first_line = lines[0].strip() if lines else ""
            detected_header = self._is_section_header(first_line)

            if detected_header:
                flush_current_chunk()
                current_section = detected_header
                current_type = CANONICAL_SECTION_MAP.get(detected_header, "general")
                current_role = None
                current_company = None

                # If block contained more than the header line, add remaining lines
                rem = "\n".join(lines[1:]).strip()
                if rem:
                    section_buf.append(rem)
            else:
                # In experience sections, detect role/company transitions:
                if current_type == "experience":
                    role_match = re.search(
                        r"(?i)\b(engineer|developer|architect|lead|manager|scientist|consultant|analyst|designer|intern)\b",
                        first_line,
                    )
                    has_date_or_company = bool(
                        re.search(
                            r"(?i)(?:\b(?:at|@)\b|\d{4}|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\b|present)",
                            first_line,
                        )
                    )
                    if role_match and has_date_or_company:
                        # Flush previous role's chunk if accumulated
                        flush_current_chunk()
                        current_role = first_line

                section_buf.append(b)

        flush_current_chunk()

        # If structuring resulted in 0 chunks (empty), create a fallback general chunk if content exists
        if not chunks and blocks:
            all_text = "\n\n".join(blocks).strip()
            if all_text:
                chunks.append(
                    ParsedSectionChunk(
                        section_title="GENERAL",
                        section_type="general",
                        content=all_text,
                        raw_content=all_text,
                        block_index=0,
                    )
                )

        return chunks

    def parse_document(
        self,
        filepath_or_name: str,
        content: Optional[str] = None,
        stream_bytes: Optional[bytes] = None,
    ) -> List[ParsedSectionChunk]:
        """
        High-level dispatcher selecting the optimal layout parser for the document.
        Ensures fallback to text parsing if layout block extraction fails.
        """
        ext = Path(filepath_or_name).suffix.lower()

        if ext == ".pdf":
            if stream_bytes:
                chunks = self.parse_pdf(stream_bytes)
            else:
                chunks = self.parse_pdf(filepath_or_name)
            if chunks:
                return chunks

        elif ext == ".docx":
            if stream_bytes:
                chunks = self.parse_docx(stream_bytes)
            else:
                chunks = self.parse_docx(filepath_or_name)
            if chunks:
                return chunks

        # Fallback to text parsing
        if content:
            return self.parse_text(content)
        if stream_bytes:
            try:
                txt = stream_bytes.decode("utf-8", errors="ignore")
                return self.parse_text(txt)
            except Exception:
                pass

        return []
